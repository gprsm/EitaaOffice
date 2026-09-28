"""Unit and integration tests for the ReportingStore and service persistence."""

from __future__ import annotations

from datetime import date, datetime, timezone
from pathlib import Path
import pytest

from eitaa_bridge.reporting import (
    DialogWatchConfig,
    EventCandidate,
    Fact,
    FactValueKind,
    FilledForm,
    IndexDecision,
    MessageIntent,
    OccasionClass,
    ProgramId,
    ProgramKind,
    ReportedEvent,
    ReportingService,
    ReportingStore,
    UnitScope,
    ValueSource,
)


@pytest.fixture
def store(tmp_path: Path) -> ReportingStore:
    db_file = tmp_path / "test_reporting.db"
    return ReportingStore(db_file)


# --------------------------------------------------------------------------
# Schema & Target Configuration Tests
# --------------------------------------------------------------------------

class TestStoreInitialization:
    def test_seeds_default_targets(self, store: ReportingStore) -> None:
        targets = store.list_targets()
        assert len(targets) >= 2
        labels = [t.label for t in targets]
        assert "مدیریت امور فرهنگی دادگستری" in labels
        assert "رابطان فرهنگی دادگستری" in labels
        assert all(t.enabled for t in targets)

    def test_save_and_delete_target(self, store: ReportingStore) -> None:
        new_target = DialogWatchConfig(
            label="کانال اطلاع‌رسانی دادگاه تجدیدنظر",
            match_keywords=("تجدیدنظر", "فرهنگی"),
            enabled=True,
        )
        store.save_target(new_target)
        targets = store.list_targets()
        assert any(t.label == "کانال اطلاع‌رسانی دادگاه تجدیدنظر" for t in targets)

        deleted = store.delete_target("کانال اطلاع‌رسانی دادگاه تجدیدنظر")
        assert deleted is True
        targets_after = store.list_targets()
        assert not any(t.label == "کانال اطلاع‌رسانی دادگاه تجدیدنظر" for t in targets_after)

    def test_config_get_and_set(self, store: ReportingStore) -> None:
        cfg = store.get_config()
        assert cfg.get("report_period") == "۱۴۰۵"
        assert cfg.get("province_name") == "مازندران"

        store.set_config({"province_name": "سمنان", "report_period": "۱۴۰۶"})
        updated = store.get_config()
        assert updated["province_name"] == "سمنان"
        assert updated["report_period"] == "۱۴۰۶"


# --------------------------------------------------------------------------
# Index Decisions Logging Tests
# --------------------------------------------------------------------------

class TestIndexDecisions:
    def test_record_and_list_decisions(self, store: ReportingStore) -> None:
        decision = IndexDecision(
            message_ref="msg-101",
            dialog_label="رابطان فرهنگی دادگستری",
            intent=MessageIntent.EVENT_REPORT,
            score=4.5,
            matched_programs=("ceremony",),
            matched_execution=("برگزار شد",),
            unit_hints=("دادگستری",),
            attendee_count=50,
            occasion_class=OccasionClass.RELIGIOUS,
        )
        store.record_decision(decision)

        all_decisions = store.list_decisions()
        assert len(all_decisions) == 1
        d = all_decisions[0]
        assert d["message_ref"] == "msg-101"
        assert d["intent"] == MessageIntent.EVENT_REPORT
        assert d["score"] == 4.5
        assert d["matched_programs"] == ["ceremony"]
        assert d["attendee_count"] == 50

        # Filter by intent
        event_reports = store.list_decisions(intent=MessageIntent.EVENT_REPORT)
        assert len(event_reports) == 1
        promotional = store.list_decisions(intent=MessageIntent.PROMOTIONAL)
        assert len(promotional) == 0


# --------------------------------------------------------------------------
# Event Candidates & Human Review Queue Tests
# --------------------------------------------------------------------------

class TestEventCandidatesReviewQueue:
    def test_save_candidate_and_dedup_pending(self, store: ReportingStore) -> None:
        cand = EventCandidate(
            event_key="evt-key-1",
            suggested_kinds=(ProgramKind.CEREMONY,),
            extracted_attendees=85,
            extracted_date="1405/02/10",
            occasion_class=OccasionClass.RELIGIOUS,
            is_ashura_pilgrimage=False,
            sender_hint="رابط فرهنگی سوادکوه",
            source_message_refs=("msg-102",),
            confidence_note="Past-tense execution detected",
        )
        cand_id1 = store.save_candidate(cand, dialog_label="رابطان فرهنگی دادگستری")
        assert cand_id1.startswith("cand-")

        # Duplicate pending insertion should return existing candidate ID
        cand_id2 = store.save_candidate(cand, dialog_label="رابطان فرهنگی دادگستری")
        assert cand_id1 == cand_id2

        pending = store.list_candidates(status="pending")
        assert len(pending) == 1
        assert pending[0]["extracted_attendees"] == 85
        assert pending[0]["dialog_label"] == "رابطان فرهنگی دادگستری"

    def test_approve_and_reject_candidate(self, store: ReportingStore) -> None:
        cand1 = EventCandidate(
            event_key="evt-key-a",
            suggested_kinds=(ProgramKind.TRIP,),
            extracted_attendees=30,
            source_message_refs=("msg-103",),
        )
        cand2 = EventCandidate(
            event_key="evt-key-b",
            suggested_kinds=(ProgramKind.CONTEST,),
            extracted_attendees=40,
            source_message_refs=("msg-104",),
        )
        id1 = store.save_candidate(cand1)
        id2 = store.save_candidate(cand2)

        store.approve_candidate(id1, event_id="real-evt-1", reviewed_by="محمدی")
        store.reject_candidate(id2, reviewed_by="محمدی", reason="اطلاعیه است رویداد نیست")

        pending = store.list_candidates(status="pending")
        assert len(pending) == 0

        approved = store.list_candidates(status="approved")
        assert len(approved) == 1
        assert approved[0]["promoted_event_id"] == "real-evt-1"
        assert approved[0]["reviewed_by"] == "محمدی"

        rejected = store.list_candidates(status="rejected")
        assert len(rejected) == 1
        assert rejected[0]["rejection_reason"] == "اطلاعیه است رویداد نیست"


# --------------------------------------------------------------------------
# Reported Events and Facts Persistence Tests
# --------------------------------------------------------------------------

class TestReportedEventsStore:
    def test_save_get_and_delete_event_with_facts(self, store: ReportingStore) -> None:
        event = ReportedEvent(
            event_id="evt-ceremony-1",
            program_kinds=(ProgramKind.CEREMONY,),
            occurred_on=date(2026, 6, 15),
            unit=UnitScope.JUDICIAL_DOMAIN,
            unit_name="دادگستری بابل",
            occasion="ولادت پیامبر (ص)",
            occasion_class=OccasionClass.RELIGIOUS,
            official_present=True,
            is_standalone_titled=True,
            notes="مراسم با حضور مسئولان قضایی برگزار گردید",
            created_by="مسئول ستاد",
        )
        fact1 = Fact(
            metric="attendees",
            value=120.0,
            value_kind=FactValueKind.VERIFIED,
            source=ValueSource.EITAA,
            evidence_refs=("msg-201",),
            created_by="مسئول ستاد",
            note="تأیید بر اساس تصویر ارسالی",
        )
        fact2 = Fact(
            metric="ceremony_count",
            value=1.0,
            value_kind=FactValueKind.OBSERVED,
            source=ValueSource.MANUAL,
        )
        event.add_fact(fact1)
        event.add_fact(fact2)

        store.save_event(event)

        loaded = store.get_event("evt-ceremony-1")
        assert loaded is not None
        assert loaded.event_id == "evt-ceremony-1"
        assert loaded.unit_name == "دادگستری بابل"
        assert loaded.official_present is True
        assert len(loaded.facts) == 2

        attendees_fact = next(f for f in loaded.facts if f.metric == "attendees")
        assert attendees_fact.value == 120.0
        assert attendees_fact.value_kind == FactValueKind.VERIFIED
        assert attendees_fact.source == ValueSource.EITAA
        assert "msg-201" in attendees_fact.evidence_refs

        # Delete event
        deleted = store.delete_event("evt-ceremony-1")
        assert deleted is True
        assert store.get_event("evt-ceremony-1") is None


# --------------------------------------------------------------------------
# Filled Forms Persistence Tests
# --------------------------------------------------------------------------

class TestFilledFormsStore:
    def test_save_and_retrieve_filled_form(self, store: ReportingStore) -> None:
        form = FilledForm(program_id=ProgramId.CEREMONIES)
        form.set("ceremony_count", 12)
        form.set("audience_count", 850)
        form.set("province", "مازندران")
        store.save_filled_form(form, updated_by="اپراتور")

        retrieved = store.get_filled_form(ProgramId.CEREMONIES)
        assert retrieved is not None
        assert retrieved.answers["ceremony_count"].value == 12
        assert retrieved.answers["province"].value == "مازندران"

        all_forms = store.list_filled_forms()
        assert ProgramId.CEREMONIES.value in all_forms


# --------------------------------------------------------------------------
# Service Integration Tests
# --------------------------------------------------------------------------

class TestServicePersistenceIntegration:
    def test_service_scan_and_record(self, store: ReportingStore) -> None:
        service = ReportingService(store=store)
        messages = [
            ("msg-501", "رابطان فرهنگی دادگستری", "مراسم عزاداری با حضور ۱۰۰ نفر در دادگستری ساری برگزار شد"),
            ("msg-502", "مدیریت امور فرهنگی دادگستری", "جشنواره فرهنگی در روزهای آینده برگزار میشود"),
        ]

        result = service.scan_and_record_texts(messages, dialog_label="رابطان فرهنگی دادگستری")
        assert result.scanned == 2
        assert result.event_reports == 1
        assert result.informational == 1

        candidates = service.list_review_candidates(status="pending")
        assert len(candidates) == 1
        cand_id = candidates[0]["candidate_id"]

        # Promote candidate to event
        event = service.approve_candidate_to_event(
            cand_id,
            event_id="promoted-evt-1",
            occurred_on=date(2026, 5, 20),
            unit_name="دادگستری ساری",
            official_present=True,
            staff_member="احمدی",
        )
        assert event is not None
        assert event.event_id == "promoted-evt-1"
        assert event.unit_name == "دادگستری ساری"

        # Verify candidate is no longer pending
        pending_after = service.list_review_candidates(status="pending")
        assert len(pending_after) == 0

        # Verify event is in store
        events = service.get_stored_events()
        assert len(events) == 1
        assert events[0].event_id == "promoted-evt-1"

    def test_service_export_with_audit(self, store: ReportingStore, tmp_path: Path) -> None:
        service = ReportingService(store=store)
        event = ReportedEvent(
            event_id="evt-export-1",
            program_kinds=(ProgramKind.CEREMONY,),
            occurred_on=date(2026, 5, 1),
            unit=UnitScope.PROVINCIAL_HQ,
        )
        store.save_event(event)

        out_file = tmp_path / "audit_report.xlsx"
        exported_path = service.export_with_audit(
            destination=out_file,
            province_name="مازندران",
            report_period="۱۴۰۵",
            allow_unresolved_star=True,
            exported_by="کاربر آزمایشی",
        )
        assert exported_path.exists()

        exports = store.list_exports()
        assert len(exports) == 1
        assert exports[0]["province_name"] == "مازندران"
        assert exports[0]["report_period"] == "۱۴۰۵"
        assert exports[0]["exported_by"] == "کاربر آزمایشی"
        assert len(exports[0]["file_sha256"]) == 64


# --------------------------------------------------------------------------
# Questionnaire definitions projection & ranged export (F-095)
# --------------------------------------------------------------------------

class TestFormDefinitionsProjection:
    def test_all_official_forms_projected(self, store: ReportingStore) -> None:
        definitions = store.list_form_definitions()
        codes = {d["program_code"] for d in definitions}
        assert codes == {"80401", "80402", "80403", "80501", "80406", "80601", "80202", "80000"}

    def test_trip_definition_carries_policy_fields(self, store: ReportingStore) -> None:
        trip = store.get_form_definition("80401")
        assert trip is not None
        assert trip["program_id"] == "trip"
        assert trip["title"].startswith("80401")
        assert "اردو" in trip["operational_description"]
        assert "دستورالعمل" in trip["policy_framework"]
        import json as _json
        criteria = _json.loads(trip["monitoring_criteria_json"])
        assert isinstance(criteria, list) and len(criteria) >= 3
        questions = _json.loads(trip["questions_json"])
        auto_metrics = {q["auto_from"] for q in questions if q["auto_from"]}
        assert {"trip_count", "attendees", "insurance_count", "vehicle_count"} <= auto_metrics

    def test_visit_worksheet_fields_projected(self, store: ReportingStore) -> None:
        """فیلدهای کاربرگ بازدید استانی در تعاریف فرم‌ها تصویر شده باشند (F-096)."""
        import json as _json

        def metrics(code: str) -> set[str]:
            row = store.get_form_definition(code)
            assert row is not None, code
            return {q["auto_from"] for q in _json.loads(row["questions_json"]) if q["auto_from"]}

        # ردیف ۱: مسابقه موضوع فرزندآوری (ماده ۲۸) و تعداد برندگان
        contest = metrics("80402")
        assert {"contest_childbearing", "winners_count"} <= contest
        # ردیف ۵: وضعیت‌سنجی نمازخانه، زیباسازی، بازسازی عتبات، خادمیاری، کمک به غزه
        prayer = metrics("80501")
        assert {
            "mosque_condition_count", "beautification_count", "atrat_reconstruction_count",
            "khademiari_mosque_count", "khademiari_honoree_count", "gaza_aid_count",
        } <= prayer
        # ردیف تکریم: دسته‌بندی پنج‌گانه
        honor = metrics("80406")
        assert {
            "honor_legends_count", "honor_retirees_count", "honor_veterans_count",
            "honor_prayer_servants_count", "honor_special_count",
        } <= honor
        # منشور: شورای فرهنگی، رونمایی، گزارش اقدامات
        charter = metrics("80202")
        assert {"council_session_count", "unveiling_count", "action_report_count"} <= charter

    def test_projection_refreshes_on_reinit(self, store: ReportingStore, tmp_path: Path) -> None:
        assert store.get_form_definition("80406") is not None
        # Reopening the same DB re-upserts definitions without duplicating.
        reopened = ReportingStore(store.db_path)
        definitions = reopened.list_form_definitions()
        assert len(definitions) == 8


class TestCanonicalMandatesSeed:
    def test_visit_worksheet_mandates_seeded(self, store: ReportingStore) -> None:
        """مستندات ابلاغی دقیق کاربرگ بازدید استانی بن‌مایه شده باشند (F-096)."""
        by_id = {m.mandate_id: m for m in store.list_mandates()}
        trip = by_id["mnd-80401-trip-circular"]
        assert trip.number == "28/2/1405"
        assert trip.issued_on == "۱۴۰۵/۰۲/۲۸"

        contest = by_id["mnd-80402-contest-guideline"]
        assert "قرآن و عترت" in contest.title
        assert contest.number == "27/2/1405"
        # ثبت تذکر تعارض کد چاپ‌شده سند (Q-IR-001)
        assert "Q-IR-001" in contest.notes

        # اسناد موضوعات تکمیلی کاربرگ بازدید با کد ارجاعِ بخش
        assert by_id["mnd-counseling-policy"].program_code == "counseling"
        assert by_id["mnd-education-services-policy"].program_code == "education_services"
        assert by_id["mnd-content-production-duties"].program_code == "content_production"
        assert by_id["mnd-training-courses-plan"].program_code == "training_courses"
        assert by_id["mnd-external-collaboration-plan"].program_code == "external_collaboration"

    def test_self_healing_seed_updates_existing_rows(self, store: ReportingStore) -> None:
        """اصلاح متن بن‌مایه در کد، پس از بازگشایی پایگاه منتشر می‌شود."""
        store.db_path  # ensure initialized
        from eitaa_bridge.reporting.plans import CANONICAL_MANDATES
        assert len(CANONICAL_MANDATES) >= 15
        reopened = ReportingStore(store.db_path)
        mandates = {m.mandate_id: m for m in reopened.list_mandates()}
        assert len(mandates) >= len(CANONICAL_MANDATES)
        # متن اصلاح‌شده جایگزین نسخهٔ قدیمی شده است
        assert mandates["mnd-80202-charter-law"].number == "5000/111979"


class TestRangedExportWithAudit:
    def _make_event(self, event_id: str, occurred: date, kind: ProgramKind, attendees: int) -> ReportedEvent:
        event = ReportedEvent(
            event_id=event_id,
            program_kinds=(kind,),
            occurred_on=occurred,
            unit=UnitScope.PROVINCIAL_HQ,
        )
        event.add_fact(
            Fact(
                metric="attendees",
                value=float(attendees),
                value_kind=FactValueKind.REPORTED_BY_UNIT,
                source=ValueSource.EITAA,
            )
        )
        return event

    def test_ranged_export_counts_only_window(self, store: ReportingStore, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        service = ReportingService(store=store)
        store.save_event(self._make_event("evt-r1", date(2026, 7, 25), ProgramKind.TRIP, 10))
        store.save_event(self._make_event("evt-r2", date(2026, 8, 15), ProgramKind.TRIP, 20))
        store.save_event(self._make_event("evt-r3", date(2026, 8, 28), ProgramKind.TRIP, 40))

        out_file = tmp_path / "ranged_report.xlsx"
        exported_path = service.export_with_audit(
            destination=out_file,
            province_name="مازندران",
            report_period="تیر تا مرداد",
            allow_unresolved_star=True,
            date_from="2026-07-01",
            date_to="2026-08-20",
        )
        assert exported_path.exists()

        workbook = load_workbook(exported_path)
        # C5 = تعداد اردو: فقط evt-r1 و evt-r2 داخل بازه‌اند (evt-r3 خارج).
        assert workbook["اردو"]["C5"].value == 2
        # D5 = تعداد شرکت‌کنندگان داخل بازه.
        assert workbook["اردو"]["D5"].value == 30

        exports = store.list_exports()
        assert exports and exports[0]["total_events"] == 2

    def test_window_without_upper_bound(self, store: ReportingStore, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        service = ReportingService(store=store)
        store.save_event(self._make_event("evt-w1", date(2026, 6, 1), ProgramKind.TRIP, 5))
        store.save_event(self._make_event("evt-w2", date(2026, 8, 1), ProgramKind.TRIP, 7))

        out_file = tmp_path / "open_ranged.xlsx"
        exported_path = service.export_with_audit(
            destination=out_file,
            allow_unresolved_star=True,
            date_from="2026-07-01",
        )
        workbook = load_workbook(exported_path)
        assert workbook["اردو"]["C5"].value == 1
        assert workbook["اردو"]["D5"].value == 7

