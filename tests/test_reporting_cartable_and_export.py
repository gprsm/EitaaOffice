import json
import tempfile
from datetime import date, datetime
from pathlib import Path
import openpyxl
import pytest

from eitaa_bridge.reporting import (
    Fact,
    FactValueKind,
    UnitScope,
    ProgramId,
    ReportedEvent,
    ValueSource,
    ReportingService,
)
from eitaa_bridge.reporting.store import ReportingStore
from eitaa_bridge.reporting.excel_export import export_events_list_to_excel
from eitaa_bridge.application.reporting_api_v3 import dispatch


@pytest.fixture
def store(tmp_path: Path) -> ReportingStore:
    db_file = tmp_path / "test_reporting.db"
    return ReportingStore(db_file)


def _sample_event(event_id: str, *, program: str = "trip", macro_ind: str = "برگزاری اردوهای فرهنگی زیارتی برای کارکنان و خانواده") -> ReportedEvent:
    now = datetime(2026, 7, 10, 10, 0, 0)
    return ReportedEvent(
        event_id=event_id,
        program_kinds=(ProgramId(program),),
        occurred_on=date(2026, 7, 10),
        unit=UnitScope.PROVINCIAL_HQ,
        unit_name="ستاد دادگستری کل استان",
        occasion="اردوی خانوادگی پرسنل مشهد مقدس",
        official_present=True,
        macro_category="workbook_approved",
        custom_topic="",
        macro_indicator=macro_ind,
        micro_indicators=["اردوی سه روزه مشهد", "همایش فرهنگی خانواده"],
        notes="گزارش نهایی و مستندات اتوبوس‌ها ضمیمه گردید.",
        facts=(
            Fact(
                metric="attendees",
                value=85.0,
                value_kind=FactValueKind.OBSERVED,
                unit_of_measure="count",
                scope="province",
                evidence_refs=("msg:123",),
                source=ValueSource.EITAA,
                created_at=now,
            ),
        ),
        created_at=now,
        created_by="operator1",
        registered_by_id="user_123",
        registered_by_name="کارشناس فرهنگی",
        review_status="needs_review",
    )


def test_update_and_get_event_file(store: ReportingStore) -> None:
    ev = _sample_event("evt_test_01")
    store.save_event(ev)

    # Verify get_event_file loads facts and deserializes fields
    f = store.get_event_file("evt_test_01")
    assert f is not None
    assert f["event_id"] == "evt_test_01"
    assert f["micro_indicators"] == ["اردوی سه روزه مشهد", "همایش فرهنگی خانواده"]
    assert f["program_kinds"] == ["trip"]
    assert len(f["facts"]) == 1
    assert f["facts"][0]["metric"] == "attendees"
    assert f["facts"][0]["value"] == 85.0

    # Test update_event_fields with program_kinds and micro_indicators
    etag = f["etag"]
    new_etag = store.update_event_fields(
        "evt_test_01",
        actor="reviewer1",
        expected_etag=etag,
        fields={
            "occasion": "اردوی اصلاح‌شده پرسنل",
            "micro_indicators": ["اردوی زیارتی", "جلسه اخلاق"],
            "notes": "اصلاحات انجام شد.",
            "program_kinds": ["trip", "contest"],
        },
    )
    assert new_etag != etag

    # Verify updated fields
    f2 = store.get_event_file("evt_test_01")
    assert f2 is not None
    assert f2["occasion"] == "اردوی اصلاح‌شده پرسنل"
    assert f2["micro_indicators"] == ["اردوی زیارتی", "جلسه اخلاق"]
    assert f2["program_kinds"] == ["trip", "contest"]
    assert f2["notes"] == "اصلاحات انجام شد."


def test_list_events_filtering(store: ReportingStore) -> None:
    ev1 = _sample_event("evt_01", program="trip", macro_ind="برگزاری اردوهای فرهنگی زیارتی برای کارکنان و خانواده")
    ev2 = _sample_event("evt_02", program="prayer", macro_ind="ترویج و توسعه فرهنگ اقامه نماز")
    store.save_event(ev1)
    store.save_event(ev2)

    # Filter by macro_indicator
    status, res, _ = dispatch(
        store,
        "reviewer1",
        ["admin"],
        "GET",
        "/api/v3/reporting/events",
        query={"macro_indicator": "ترویج و توسعه فرهنگ اقامه نماز"},
    )
    assert status == 200
    events = res["events"]
    assert len(events) == 1
    assert events[0]["event_id"] == "evt_02"

    # Filter by search
    status, res, _ = dispatch(
        store,
        "reviewer1",
        ["admin"],
        "GET",
        "/api/v3/reporting/events",
        query={"search": "مشهد"},
    )
    assert status == 200
    assert len(res["events"]) == 2

    # Filter by micro_indicator
    status, res, _ = dispatch(
        store,
        "reviewer1",
        ["admin"],
        "GET",
        "/api/v3/reporting/events",
        query={"micro_indicator": "همایش"},
    )
    assert status == 200
    assert len(res["events"]) == 2


def test_export_filtered_events_excel(store: ReportingStore, tmp_path: Path) -> None:
    ev = _sample_event("evt_export_01")
    store.save_event(ev)

    status, res, _ = dispatch(
        store,
        "reviewer1",
        ["admin"],
        "POST",
        "/api/v3/reporting/events/export-excel",
        payload={"search": "مشهد", "review_status": "needs_review"},
    )
    assert status == 200
    assert res["count"] == 1
    export_path = Path(res["path"])
    assert export_path.exists()

    wb = openpyxl.load_workbook(export_path)
    assert "رویدادهای فرهنگی" in wb.sheetnames
    ws = wb["رویدادهای فرهنگی"]
    assert ws["B3"].value == "evt_export_01"
    assert "مشهد" in str(ws["K3"].value)
    wb.close()


def test_single_program_export_unified_report(store: ReportingStore, tmp_path: Path) -> None:
    # Look for existing template
    template_candidates = list(Path(".").glob("*.xlsx"))
    template_file = next((f for f in template_candidates if "1405" in f.name and not f.name.startswith("~")), None)
    if not template_file:
        pytest.skip("Template 1405.xlsx not found in root")

    dest_file = tmp_path / "single_export_test.xlsx"
    ev = _sample_event("evt_trip_01", program="trip")

    service = ReportingService(store=store)
    service.export_with_audit(
        events=[ev],
        destination=dest_file,
        report_period="۱۴۰۵",
        province_name="مازندران",
        allow_unresolved_star=True,
        template_root=template_file.parent,
        single_program="trip",
    )

    assert dest_file.exists()
    wb = openpyxl.load_workbook(dest_file)
    # Target sheet "اردو" must be preserved, while other sheets like "مسابقات", "منشور" should be deleted
    assert "اردو" in wb.sheetnames
    assert "مسابقات" not in wb.sheetnames
    assert "منشور" not in wb.sheetnames
    wb.close()


def test_single_program_export_ignores_unresolved_stars_in_other_sheets(store: ReportingStore, tmp_path: Path) -> None:
    template_candidates = list(Path(".").glob("*.xlsx"))
    template_file = next((f for f in template_candidates if "1405" in f.name and not f.name.startswith("~")), None)
    if not template_file:
        pytest.skip("Template 1405.xlsx not found in root")

    dest_file = tmp_path / "single_export_star_test.xlsx"
    ev = _sample_event("evt_charter_02", program="charter")

    from eitaa_bridge.reporting import FilledForm
    charter_form = FilledForm(ProgramId.CHARTER)
    charter_form.set("correspondence", 1)
    charter_form.set("boards_installed", 1)
    charter_form.set("republish", 1)
    store.save_filled_form(charter_form)

    service = ReportingService(store=store)
    # allow_unresolved_star is False; since charter ("منشور") is filled, missing stars in other unexported sheets should not block!
    service.export_with_audit(
        events=[ev],
        destination=dest_file,
        report_period="۱۴۰۵",
        province_name="مازندران",
        allow_unresolved_star=False,
        template_root=template_file.parent,
        single_program="charter",
    )
    assert dest_file.exists()


def test_list_events_safe_against_numeric_catalog_codes(store: ReportingStore) -> None:
    # Directly insert an event with numeric catalog code in program_kinds_json
    with store._connect() as conn:
        conn.execute(
            """
            INSERT INTO reported_events (
                event_id, occurred_on, unit, unit_name, occasion,
                macro_category, custom_topic, macro_indicator,
                micro_indicators_json, program_kinds_json, created_at, created_by,
                review_status, registered_by_id, registered_by_name, version
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "evt_num_01", "2026-07-10", "provincial_hq", "ستاد", "مناسبت تستی",
                "workbook_approved", "", "شاخص کلی",
                json.dumps(["شاخص جزئی"]), json.dumps(["80401", "custom_code_99"]),
                "2026-07-10T10:00:00Z", "op", "needs_review", "user", "کارشناس", 1
            )
        )
    # list_events must not crash
    events = store.list_events()
    assert any(e.event_id == "evt_num_01" for e in events)


def test_honor_form_prefill_and_aggregation_with_metric_aliases() -> None:
    from eitaa_bridge.reporting.model import (
        ReportedEvent,
        ProgramKind,
        UnitScope,
        Fact,
        FactValueKind,
        ValueSource,
        ProgramId,
    )
    from eitaa_bridge.reporting.forms import HONOR_FORM, prefill_form
    from eitaa_bridge.reporting.aggregate import aggregate_program_report
    from datetime import date

    ev = ReportedEvent(
        event_id="evt-test-honor-01",
        program_kinds=(ProgramKind.HONOR,),
        occurred_on=date(2026, 10, 19),
        unit=UnitScope.JUDICIAL_DOMAIN,
        unit_name="سوادکوه - دادگستری",
        occasion="تقدیر از همکار روشن دل",
        official_present=True,
        is_standalone_titled=True,
        facts=[
            Fact(
                metric="honored_count",
                value=1.0,
                value_kind=FactValueKind.VERIFIED,
                unit_of_measure="نفر",
                source=ValueSource.MANUAL,
            )
        ],
    )

    # 1. prefill_form must map honored_count to honorees and ceremony_count to 1
    filled = prefill_form(HONOR_FORM, [ev])
    assert filled.answers["honorees"].value == 1
    assert filled.answers["ceremony_count"].value == 1

    # 2. aggregate_program_report must sum into counts['honorees'] and breakdown_by_unit
    rep = aggregate_program_report(ProgramId.HONOR, [ev])
    assert rep.counts["honor"] == 1
    assert rep.counts["honorees"] == 1
    assert rep.breakdown_by_unit["سوادکوه - دادگستری"]["honorees"] == 1
    assert rep.breakdown_by_unit["سوادکوه - دادگستری"]["events"] == 1

