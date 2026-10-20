"""Excel export: project a unified report onto a *copy* of the 1405 workbook.

The original workbook ``SRC-IR-001`` is a read-only source template; export
always writes to a fresh copy (ADR-53). Cell coordinates mirror the template
exactly so the human final pass edits a familiar layout. Star-marked cells
must receive a value; anything missing is reported as a blocker, not guessed.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping, Sequence

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .aggregate import ProgramReport
from .forms import QuestionnaireDefinition
from .model import FactValueKind, ProgramId, ReportedEvent
from .rules import RULES_VERSION

TEMPLATE_VERSION = "workbook-1405-template-v1"

TEMPLATE_FILENAME = "-_امور_فرهنگي_.فرمت_گزارش__برنامه_هاي_متناظر_استاني_1405.xlsx"
ASHURA_ANNEX_SHEET_NAME = "ضمیمه زیارت عاشورا"


class ReportingExportError(Exception):
    code = "reporting_export_error"


class TemplateNotFoundError(ReportingExportError):
    code = "template_not_found"


class UnresolvedStarCellsError(ReportingExportError):
    code = "unresolved_star_cells"

    def __init__(self, blockers: dict[str, list[str]]) -> None:
        super().__init__(f"Unresolved star cells: {blockers}")
        self.blockers = blockers


def find_template(root: Path | None = None) -> Path:
    base = root or Path(__file__).resolve().parents[3]
    candidate = base / TEMPLATE_FILENAME
    if candidate.exists():
        return candidate
    raise TemplateNotFoundError(f"Workbook template not found at {candidate}")


def _ensure_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _answers_by_key(form) -> Mapping[str, Any]:
    return {key: answer.value for key, answer in form.answers.items()}


def export_unified_report(
    unified: Mapping[str, ProgramReport],
    filled_forms: Mapping[str, Any],
    destination: Path,
    *,
    template_root: Path | None = None,
    province_name: str = "",
    report_period: str = "۱۴۰۵",
    allow_unresolved_star: bool = False,
    ashura_events: Sequence[ReportedEvent] = (),
    traceability: Mapping[str, Any] | None = None,
    single_program: str | None = None,
) -> Path:
    """Write one aggregate row per program sheet onto a copy of the template.

    ``filled_forms`` maps program id -> ``FilledForm`` with human-reviewed
    answers; those take precedence, otherwise aggregates fill counts.
    """

    template = find_template(template_root)
    destination = destination.expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(template)

    blockers: dict[str, list[str]] = {}

    def _cell(sheet: str, coordinate: str) -> Any:
        return workbook[sheet][coordinate]

    def _write(sheet: str, coordinate: str, value: Any) -> None:
        _cell(sheet, coordinate).value = value

    def _put(sheet: str, coordinate: str, value: Any, *, star: bool, label: str) -> None:
        if value is None or value == "":
            if star and not allow_unresolved_star:
                blockers.setdefault(sheet, []).append(f"{coordinate} ({label})")
            return
        _write(sheet, coordinate, value)

    # ---- PRG-WB-01 اردو -------------------------------------------------
    if "trip" in filled_forms or "trip" in unified:
        answers = _answers_by_key(filled_forms.get("trip")) if "trip" in filled_forms else {}
        report = unified.get("trip")
        _write("اردو", "I2", f"استان: {province_name}" if province_name else "استان:")
        _put("اردو", "C5", answers.get("trip_count", report.counts.get("trip") if report else None), star=True, label="تعداد اردو")
        _put("اردو", "D5", answers.get("attendees", report.attendees_total if report else None), star=True, label="تعداد شرکت‌کنندگان")
        _put("اردو", "F5", answers.get("announcement"), star=True, label="اطلاع‌رسانی")
        _put("اردو", "G5", answers.get("attendee_list"), star=True, label="تعداد و لیست شرکت‌کنندگان")
        _put("اردو", "K5", answers.get("reception"), star=True, label="پذیرایی")
        _put("اردو", "L5", answers.get("trip_type_staff"), star=True, label="نوع اردو — کارکنان")
        _put("اردو", "M5", answers.get("trip_type_family"), star=True, label="نوع اردو — خانواده")
        _put("اردو", "N5", answers.get("trip_type_marriage"), star=True, label="نوع اردو — فرزندآوری و ازدواج")

    # ---- PRG-WB-02 مسابقات (two tables) ----------------------------------
    if "contest" in filled_forms or "contest" in unified:
        answers = _answers_by_key(filled_forms.get("contest")) if "contest" in filled_forms else {}
        _write("مسابقات", "C2", f"استان: {province_name}" if province_name else "استان:")
        _put("مسابقات", "C6", answers.get("quran_attendees"), star=True, label="قرآنی — شرکت‌کنندگان")
        _put("مسابقات", "F6", answers.get("quran_announcement"), star=True, label="قرآنی — اطلاع‌رسانی")
        _put("مسابقات", "G6", answers.get("quran_list"), star=True, label="قرآنی — لیست شرکت‌کنندگان")
        _put("مسابقات", "I6", answers.get("quran_reception"), star=True, label="قرآنی — پذیرایی")
        _put("مسابقات", "K6", answers.get("quran_awards"), star=True, label="قرآنی — تقدیر از برگزیدگان")
        report = unified.get("contest")
        _put("مسابقات", "C12", answers.get("contest_count", report.counts.get("contest") if report else None), star=True, label="تعداد مسابقه")
        _put("مسابقات", "G12", answers.get("attendees", report.attendees_total if report else None), star=True, label="شرکت‌کنندگان")
        _put("مسابقات", "I12", answers.get("announcement"), star=True, label="اطلاع‌رسانی")
        _put("مسابقات", "J12", answers.get("list"), star=True, label="لیست شرکت‌کنندگان")
        _put("مسابقات", "N12", answers.get("awards"), star=True, label="تقدیر از برگزیدگان")
        _put("مسابقات", "P12", answers.get("audience_staff"), star=True, label="مخاطب — کارکنان")
        _put("مسابقات", "Q12", answers.get("audience_family"), star=True, label="مخاطب — خانواده")
        _put("مسابقات", "R12", answers.get("audience_women"), star=True, label="مخاطب — بانوان")
        _put("مسابقات", "S12", answers.get("audience_children"), star=True, label="مخاطب — فرزندان")

    # ---- PRG-WB-03 مراسم (with ashura annex) -----------------------------
    if "ceremonies" in filled_forms or "ceremonies" in unified:
        answers = _answers_by_key(filled_forms.get("ceremonies")) if "ceremonies" in filled_forms else {}
        report = unified.get("ceremonies")
        _write("مراسم  مذهبی", "H4", f"استان: {province_name}" if province_name else "استان:")
        _put("مراسم  مذهبی", "D8", answers.get("ceremony_national", report.counts.get("ceremony_national") if report else None), star=True, label="مراسم ملی")
        _put("مراسم  مذهبی", "E8", answers.get("ceremony_religious", report.counts.get("ceremony_religious") if report else None), star=True, label="مراسم مذهبی")
        _put("مراسم  مذهبی", "F8", answers.get("ceremony_revolutionary", report.counts.get("ceremony_revolutionary") if report else None), star=True, label="مراسم انقلابی")
        _put("مراسم  مذهبی", "G8", answers.get("attendees", report.attendees_total if report else None), star=True, label="شرکت‌کنندگان")
        _put("مراسم  مذهبی", "I8", answers.get("announcement"), star=True, label="اطلاع‌رسانی")
        _put("مراسم  مذهبی", "K8", answers.get("culture_pack"), star=True, label="بستهٔ فرهنگی")
        if ashura_events or (report and report.ashura_pilgrimage_count):
            count = answers.get("ashura_pilgrimage_count", report.ashura_pilgrimage_count if report else None)
            attendees = answers.get("ashura_pilgrimage_attendees", report.ashura_pilgrimage_attendees if report else None)
            annex = workbook.create_sheet(ASHURA_ANNEX_SHEET_NAME) if ASHURA_ANNEX_SHEET_NAME not in workbook.sheetnames else workbook[ASHURA_ANNEX_SHEET_NAME]
            annex["A1"] = "ضمیمه زیارت عاشورا (جدا از آمار اصلی مراسم — C12)"
            annex["A3"] = "تعداد"
            annex["B3"] = "تعداد شرکت‌کنندگان"
            annex["A4"] = count if count is not None else 0
            annex["B4"] = attendees if attendees is not None else 0
            annex["A6"] = "رویدادها (تفکیک حوزه/هر زیارت):"
            row = 7
            for event in ashura_events:
                annex[f"A{row}"] = event.event_id
                annex[f"B{row}"] = next((int(f.value) for f in event.facts if f.metric == "attendees"), None)
                annex[f"C{row}"] = event.unit_name or event.unit.value
                row += 1

    # ---- PRG-WB-04 اقامه نماز -------------------------------------------
    if "prayer" in filled_forms or "prayer" in unified:
        answers = _answers_by_key(filled_forms.get("prayer")) if "prayer" in filled_forms else {}
        report = unified.get("prayer")
        _write("ترویج و توسعه فرهنگ اقامه نماز", "A2", f"استان: {province_name}" if province_name else "استان:")
        _put("ترویج و توسعه فرهنگ اقامه نماز", "C6", answers.get("nominees"), star=True, label="افراد نومکلف")
        _put("ترویج و توسعه فرهنگ اقامه نماز", "D6", answers.get("attendees", report.attendees_total if report else None), star=True, label="شرکت‌کنندگان")
        _put("ترویج و توسعه فرهنگ اقامه نماز", "E6", answers.get("announcement"), star=True, label="اطلاع‌رسانی")
        _put("ترویج و توسعه فرهنگ اقامه نماز", "G6", answers.get("reception"), star=True, label="پذیرایی")
        _put("ترویج و توسعه فرهنگ اقامه نماز", "H6", answers.get("gifts"), star=True, label="هدایا و بستهٔ فرهنگی")

    # ---- PRG-WB-05 تکریم و تجلیل -----------------------------------------
    if "honor" in filled_forms or "honor" in unified:
        answers = _answers_by_key(filled_forms.get("honor")) if "honor" in filled_forms else {}
        report = unified.get("honor")
        _write("تکریم و تجلیل", "A2", f"استان: {province_name}" if province_name else "استان:")
        _put("تکریم و تجلیل", "B5", answers.get("ceremony_count", report.counts.get("honor") if report else None), star=True, label="تعداد مراسم")
        _put("تکریم و تجلیل", "C5", answers.get("honorees"), star=True, label="تعداد تقدیرشدگان")
        _put("تکریم و تجلیل", "D5", answers.get("announcement"), star=True, label="اطلاع‌رسانی")
        _put("تکریم و تجلیل", "E5", answers.get("invitee_list"), star=True, label="لیست دعوت‌شدگان")
        _put("تکریم و تجلیل", "G5", answers.get("gifts"), star=True, label="هدایا")

    # ---- PRG-WB-06 تشویق ارباب رجوع --------------------------------------
    if "customer_care" in filled_forms or "customer_care" in unified:
        answers = _answers_by_key(filled_forms.get("customer_care")) if "customer_care" in filled_forms else {}
        report = unified.get("customer_care")
        _write("تشویق ارباب رجوع", "B2", f"استان: {province_name}" if province_name else "استان:")
        _put("تشویق ارباب رجوع", "C5", answers.get("ceremony_count", report.counts.get("customer_care") if report else None), star=True, label="تعداد مراسم")
        _put("تشویق ارباب رجوع", "D5", answers.get("honorees"), star=True, label="تعداد تقدیرشدگان")
        _put("تشویق ارباب رجوع", "E5", answers.get("announcement"), star=True, label="اطلاع‌رسانی")
        _put("تشویق ارباب رجوع", "F5", answers.get("invitee_list"), star=True, label="لیست دعوت‌شدگان")
        _put("تشویق ارباب رجوع", "H5", answers.get("gifts"), star=True, label="هدایا")

    # ---- PRG-WB-07 منشور -------------------------------------------------
    if "charter" in filled_forms or "charter" in unified:
        answers = _answers_by_key(filled_forms.get("charter")) if "charter" in filled_forms else {}
        _write("منشور", "B2", f"استان: {province_name}" if province_name else "استان:")
        _put("منشور", "C5", answers.get("correspondence"), star=True, label="مکاتبات")
        _put("منشور", "D5", answers.get("boards_installed"), star=True, label="تابلو نصب‌شده")
        _put("منشور", "E5", answers.get("republish"), star=True, label="بازنشر")

    # Single program sheet filter if requested
    target_sheet_name: str | None = None
    keep_sheets: set[str] | None = None
    if single_program:
        prog_sheets = {
            "trip": "اردو",
            "80401": "اردو",
            "contest": "مسابقات",
            "80402": "مسابقات",
            "ceremonies": "مراسم  مذهبی",
            "ceremony": "مراسم  مذهبی",
            "80403": "مراسم  مذهبی",
            "prayer": "ترویج و توسعه فرهنگ اقامه نماز",
            "80501": "ترویج و توسعه فرهنگ اقامه نماز",
            "honor": "تکریم و تجلیل",
            "80406": "تکریم و تجلیل",
            "customer_care": "تشویق ارباب رجوع",
            "80601": "تشویق ارباب رجوع",
            "charter": "منشور",
            "80202": "منشور",
        }
        target_s = prog_sheets.get(str(single_program).strip().lower(), str(single_program).strip())
        target_clean = target_s.replace(" ", "")
        for sname in workbook.sheetnames:
            if sname == target_s or sname.replace(" ", "") == target_clean:
                target_sheet_name = sname
                break
        if target_sheet_name:
            keep_sheets = {target_sheet_name}
            if ("مراسم" in target_sheet_name or "ceremon" in str(single_program).lower()) and ASHURA_ANNEX_SHEET_NAME in workbook.sheetnames:
                keep_sheets.add(ASHURA_ANNEX_SHEET_NAME)
            # Only star cells in the retained target sheet should block single-sheet export
            blockers = {s: b for s, b in blockers.items() if s in keep_sheets}

    if blockers and not allow_unresolved_star:
        raise UnresolvedStarCellsError(blockers)

    # Phase 4: traceability sheet + version stamps (strategy sections 3 and 9).
    if "ردیابی" in workbook.sheetnames:
        del workbook["ردیابی"]
    trace_sheet = workbook.create_sheet("ردیابی")
    trace_sheet["A1"] = "نسخهٔ قالب"
    trace_sheet["B1"] = TEMPLATE_VERSION
    trace_sheet["A2"] = "نسخهٔ قواعد"
    trace_sheet["B2"] = RULES_VERSION
    trace_sheet["A4"] = "برنامه"
    trace_sheet["B4"] = "سنجه"
    trace_sheet["C4"] = "مقدار"
    trace_sheet["D4"] = "شناسهٔ رویدادها"
    row_cursor = 5
    for program_id, payload in sorted((traceability or {}).get("programs", {}).items()):
        trace_sheet.cell(row=row_cursor, column=1, value=program_id)
        trace_sheet.cell(row=row_cursor, column=2, value="count")
        trace_sheet.cell(row=row_cursor, column=3, value=payload.get("count", 0))
        trace_sheet.cell(row=row_cursor, column=4, value="،".join(payload.get("event_ids", [])))
        row_cursor += 1
        trace_sheet.cell(row=row_cursor, column=2, value="attendees_total")
        trace_sheet.cell(row=row_cursor, column=3, value=payload.get("attendees_total", 0))
        trace_sheet.cell(row=row_cursor, column=4, value="،".join(payload.get("attendee_event_ids", [])))
        row_cursor += 1
    for item in (traceability or {}).get("excluded", []):
        trace_sheet.cell(row=row_cursor, column=1, value="خارج‌شده")
        trace_sheet.cell(row=row_cursor, column=2, value=item.get("reason", ""))
        trace_sheet.cell(row=row_cursor, column=4, value=item.get("event_id", ""))
        row_cursor += 1

    # Remove non-target sheets if single_program was requested
    if keep_sheets:
        keep = set(keep_sheets)
        if "ردیابی" in workbook.sheetnames:
            keep.add("ردیابی")
        for sname in list(workbook.sheetnames):
            if sname not in keep:
                del workbook[sname]

    workbook.save(destination)
    return destination


def star_cell_report(
    definitions: Sequence[QuestionnaireDefinition],
    filled_forms: Mapping[str, Any],
) -> dict[str, list[str]]:
    """List unresolved star questions per program before export."""

    unresolved: dict[str, list[str]] = {}
    for definition in definitions:
        form = filled_forms.get(definition.program_id.value)
        if form is None:
            unresolved[definition.program_id.value] = list(definition.star_keys())
            continue
        missing = form.unresolved_star_keys(definition)
        if missing:
            unresolved[definition.program_id.value] = list(missing)
    return unresolved


def export_events_list_to_excel(
    events: Sequence[ReportedEvent],
    destination: Path,
    *,
    province_name: str = "مازندران",
    title: str = "گزارش رویدادها و فعالیت‌های فرهنگی",
) -> Path:
    destination = destination.expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "رویدادهای فرهنگی"
    ws.sheet_view.rightToLeft = True

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Tahoma", size=13, bold=True, color="1F4E79")
    data_font = Font(name="Tahoma", size=9)
    bold_data_font = Font(name="Tahoma", size=9, bold=True)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)
    border_side = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = f"{title} — استان {province_name}"
    title_cell.font = title_font
    title_cell.alignment = center_align
    ws.row_dimensions[1].height = 32

    headers = [
        ("ردیف", 6),
        ("شناسه رویداد", 16),
        ("سرفصل / برنامه", 28),
        ("شاخصه کلی", 28),
        ("شاخصه‌های جزئی", 30),
        ("تاریخ وقوع", 14),
        ("واحد سازمانی / حوزه قضایی", 26),
        ("ثبت‌کننده", 22),
        ("وضعیت بررسی", 16),
        ("سنجه‌ها و آمار", 30),
        ("توضیحات و خلاصه گزارش", 40),
    ]

    for col_idx, (h_title, width) in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=h_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = cell_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[2].height = 26

    status_map = {
        "approved": "تأیید مدیر",
        "needs_review": "در انتظار بررسی",
        "conflict": "تعارض / رد شده",
        "draft": "پیش‌نویس",
    }

    for row_idx, ev in enumerate(events, 3):
        def _fmt_val(v: Any) -> str:
            try:
                fv = float(v)
                return str(int(fv) if fv.is_integer() else fv)
            except (TypeError, ValueError):
                return str(v or "")

        facts_summary = "، ".join(
            f"{f.metric}: {_fmt_val(f.value)} {f.unit_of_measure or ''}".strip()
            for f in ev.facts
        )
        micros_summary = "، ".join(ev.micro_indicators) if ev.micro_indicators else "—"
        program_summary = ev.custom_topic if ev.macro_category == "other_cultural" and ev.custom_topic else (
            "، ".join(k.value if hasattr(k, "value") else str(k) for k in ev.program_kinds) or "—"
        )
        status_label = status_map.get(ev.review_status, ev.review_status)

        row_data = [
            (row_idx - 2, center_align, data_font),
            (ev.event_id, center_align, data_font),
            (program_summary, right_align, bold_data_font),
            (ev.macro_indicator or "—", right_align, data_font),
            (micros_summary, right_align, data_font),
            (ev.occurred_on.isoformat() if hasattr(ev.occurred_on, "isoformat") else str(ev.occurred_on or "—"), center_align, data_font),
            (ev.unit_name or ev.unit.value, right_align, data_font),
            (ev.registered_by_name or ev.created_by or "—", right_align, data_font),
            (status_label, center_align, bold_data_font),
            (facts_summary or "—", right_align, data_font),
            (ev.occasion or ev.notes or "—", right_align, data_font),
        ]

        ws.row_dimensions[row_idx].height = 24
        for col_idx, (val, align, font) in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = align
            cell.font = font
            cell.border = cell_border

    wb.save(destination)
    return destination
